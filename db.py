# db.py
# Работа с базой данных (aiosqlite) для бота.
# - Сохраняем даты в формате 'YYYY-MM-DD HH:MM:SS' (ISO)
# - Возвращаем пользователю и в экспорте даты в формате 'DD.MM.YYYY HH:MM'
# - Поддерживаем чтение старых данных в формате 'DD.MM.YYYY HH:MM'

import aiosqlite
from datetime import datetime, timedelta, date
import os
from typing import Optional, List, Tuple
import openpyxl
from openpyxl.utils import get_column_letter
import config

DB: Optional[aiosqlite.Connection] = None

# -------------------------
# Вспомогательные функции
# -------------------------

def _now_db_str() -> str:
    """Текущее время UTC в формате ISO для хранения в базе"""
    return datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')

def _parse_db_datetime(s: Optional[str]) -> Optional[datetime]:
    """
    Преобразует строку даты из базы в объект datetime.
    Поддерживает несколько форматов для устойчивости:
    - 'YYYY-MM-DD HH:MM:SS'
    - 'YYYY-MM-DDTHH:MM:SS'
    - 'DD.MM.YYYY HH:MM'  (старый формат)
    """
    if not s:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(s, fmt)
        except Exception:
            continue
    # Попытка удалить миллисекунды
    try:
        if '.' in s:
            base = s.split('.', 1)[0]
            return datetime.strptime(base, "%Y-%m-%d %H:%M:%S")
    except Exception:
        pass
    return None

def _to_display(dt_str: Optional[str]) -> Optional[str]:
    """Формат для отображения пользователю: 'DD.MM.YYYY HH:MM'"""
    dt = _parse_db_datetime(dt_str)
    if not dt:
        return None
    return dt.strftime("%d.%m.%Y %H:%M")

def _day_to_display(day_str: Optional[str]) -> str:
    """
    day_str приходит из SQL как 'YYYY-MM-DD'.
    Возвращаем 'DD.MM.YYYY 00:00'.
    """
    if not day_str:
        return ''
    try:
        d = datetime.strptime(day_str, "%Y-%m-%d").date()
        return d.strftime("%d.%m.%Y") + " 00:00"
    except Exception:
        return day_str

# -------------------------
# Инициализация базы
# -------------------------

async def init_db(db_path: str = 'base.db'):
    global DB
    DB = await aiosqlite.connect(db_path)
    await DB.execute("""CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tg_id INTEGER UNIQUE,
        name TEXT,
        created_at TEXT
    )""")
    await DB.execute("""CREATE TABLE IF NOT EXISTS movies (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER,
        title TEXT,
        year INTEGER,
        rating INTEGER,
        review TEXT,
        created_at TEXT,
        FOREIGN KEY(user_id) REFERENCES users(id)
    )""")
    await DB.execute("""CREATE TABLE IF NOT EXISTS admins (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tg_id INTEGER UNIQUE,
        note TEXT,
        created_at TEXT
    )""")
    await DB.commit()

# -------------------------
# Пользователи и фильмы
# -------------------------

async def add_user(tg_id: int, name: Optional[str] = None):
    cur = await DB.execute("SELECT id FROM users WHERE tg_id = ?", (tg_id,))
    row = await cur.fetchone()
    if row:
        return row[0]
    now = _now_db_str()
    cur = await DB.execute(
        "INSERT INTO users (tg_id, name, created_at) VALUES (?, ?, ?)", (tg_id, name, now)
    )
    await DB.commit()
    return cur.lastrowid

async def get_user_by_tgid(tg_id: int):
    cur = await DB.execute("SELECT id, tg_id, name, created_at FROM users WHERE tg_id = ?", (tg_id,))
    row = await cur.fetchone()
    if not row:
        return None
    uid, tg, name, created_at = row
    return (uid, tg, name, _to_display(created_at))

async def add_movie(user_id: int, title: str, year: int, rating: int, review: str):
    now = _now_db_str()
    cur = await DB.execute(
        "INSERT INTO movies (user_id, title, year, rating, review, created_at) VALUES (?, ?, ?, ?, ?, ?)",
        (user_id, title, year, rating, review, now)
    )
    await DB.commit()
    return cur.lastrowid

async def get_movies(user_id: int):
    cur = await DB.execute(
        "SELECT id, title, year, rating, review, created_at FROM movies WHERE user_id = ? ORDER BY created_at DESC",
        (user_id,)
    )
    rows = await cur.fetchall()
    result = []
    for r in rows:
        mid, title, year, rating, review, created_at = r
        result.append((mid, title, year, rating, review, _to_display(created_at)))
    return result

async def get_movie(user_id: int, movie_id: int):
    cur = await DB.execute(
        "SELECT id, title, year, rating, review, created_at FROM movies WHERE user_id = ? AND id = ?",
        (user_id, movie_id)
    )
    row = await cur.fetchone()
    if not row:
        return None
    mid, title, year, rating, review, created_at = row
    return (mid, title, year, rating, review, _to_display(created_at))

async def search_movies(user_id: int, keyword: str):
    q = f"%{keyword}%"
    cur = await DB.execute(
        "SELECT id, title, year, rating, review, created_at FROM movies WHERE user_id = ? AND (title LIKE ? OR review LIKE ?) ORDER BY created_at DESC",
        (user_id, q, q)
    )
    rows = await cur.fetchall()
    result = []
    for r in rows:
        mid, title, year, rating, review, created_at = r
        result.append((mid, title, year, rating, review, _to_display(created_at)))
    return result

async def get_stats(user_id: int):
    cur = await DB.execute("SELECT COUNT(*), AVG(rating) FROM movies WHERE user_id = ?", (user_id,))
    cnt_row = await cur.fetchone()
    total = cnt_row[0] or 0
    avg = cnt_row[1] or 0
    cur2 = await DB.execute(
        "SELECT title, rating FROM movies WHERE user_id = ? ORDER BY rating DESC, created_at ASC LIMIT 1",
        (user_id,)
    )
    top = await cur2.fetchone()
    top_title = top[0] if top else None
    top_rating = top[1] if top else None
    return {'total': total, 'avg': avg, 'top_title': top_title, 'top_rating': top_rating}

async def get_movies_count_by_day(user_id: int, days: int):
    start_date = (datetime.utcnow() - timedelta(days=days-1)).date().strftime('%Y-%m-%d')
    cur = await DB.execute(
        """SELECT date(created_at) as day, COUNT(*) FROM movies WHERE user_id = ? AND date(created_at) >= ? GROUP BY day ORDER BY day""",
        (user_id, start_date)
    )
    rows = await cur.fetchall()
    counts = {r[0]: r[1] for r in rows}
    result = []
    start_dt = (datetime.utcnow() - timedelta(days=days-1)).date()
    for i in range(days):
        d = start_dt + timedelta(days=i)
        day_key = d.strftime('%Y-%m-%d')
        display = d.strftime("%d.%m.%Y") + " 00:00"
        result.append((display, counts.get(day_key, 0)))
    return result

async def get_all_tg_users():
    cur = await DB.execute("SELECT tg_id FROM users")
    rows = await cur.fetchall()
    return [r[0] for r in rows]

# -------------------------
# Админы
# -------------------------

async def add_admin(tg_id: int, note: Optional[str] = None):
    now = _now_db_str()
    await DB.execute("INSERT OR IGNORE INTO admins (tg_id, note, created_at) VALUES (?, ?, ?)", (tg_id, note, now))
    await DB.commit()

async def remove_admin(tg_id: int):
    await DB.execute("DELETE FROM admins WHERE tg_id = ?", (tg_id,))
    await DB.commit()

async def is_admin(tg_id: int) -> bool:
    cur = await DB.execute("SELECT id FROM admins WHERE tg_id = ?", (tg_id,))
    row = await cur.fetchone()
    return True if row or tg_id in config.ADMINS else False

async def get_all_admins() -> List[int]:
    cur = await DB.execute("SELECT tg_id FROM admins")
    rows = await cur.fetchall()
    return [r[0] for r in rows]

# -------------------------
# Экспорт в Excel
# -------------------------

async def export_all_to_excel(tmp_dir: str = 'tmp', filename: Optional[str] = None) -> str:
    os.makedirs(tmp_dir, exist_ok=True)
    if filename is None:
        filename = f"export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = os.path.join(tmp_dir, filename)

    cur = await DB.execute("SELECT id, tg_id, name, created_at FROM users")
    users_rows = await cur.fetchall()
    users_by_id = {r[0]: r for r in users_rows}
    users_by_tgid = {r[1]: r for r in users_rows if r[1] is not None}

    cur2 = await DB.execute("SELECT id, user_id, title, year, rating, review, created_at FROM movies ORDER BY user_id, created_at")
    movies_rows = await cur2.fetchall()

    rows = []
    for m in movies_rows:
        mid, m_user_id, title, year, rating, review, m_created = m
        u = users_by_id.get(m_user_id)
        if not u:
            u = users_by_tgid.get(m_user_id)
        if u:
            uid, tg_id, name, u_created = u
            u_created_disp = _to_display(u_created)
        else:
            uid = None
            tg_id = None
            name = None
            u_created_disp = None
        m_created_disp = _to_display(m_created)
        rows.append({
            'user_id': uid,
            'user_tg_id': tg_id,
            'user_name': name,
            'user_created_at': u_created_disp,
            'movie_id': mid,
            'title': title,
            'year': year,
            'rating': rating,
            'review': review,
            'movie_created_at': m_created_disp
        })

    def sort_key(r):
        return (r['user_tg_id'] if r['user_tg_id'] is not None else 10**18, r['movie_created_at'] or '')
    rows.sort(key=sort_key)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "users_movies"

    headers = ['user_id (users.id)', 'user_tg_id (users.tg_id)', 'user_name', 'user_created_at',
               'movie_id', 'title', 'year', 'rating', 'review', 'movie_created_at']
    ws.append(headers)

    for r in rows:
        ws.append([
            r['user_id'],
            r['user_tg_id'],
            r['user_name'],
            r['user_created_at'],
            r['movie_id'],
            r['title'],
            r['year'],
            r['rating'],
            r['review'],
            r['movie_created_at']
        ])

    for i, col in enumerate(ws.columns, 1):
        max_len = 0
        for cell in col:
            try:
                v = cell.value or ""
                l = len(str(v))
                if l > max_len:
                    max_len = l
            except:
                pass
        adjusted_width = min(max(10, max_len + 2), 60)
        ws.column_dimensions[get_column_letter(i)].width = adjusted_width

    wb.save(path)
    return path
